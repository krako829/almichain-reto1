#!/usr/bin/env python3
"""Almi-Chain · Portada, cabecera y pie de la memoria.
  python3 portada.py <carpeta de trabajo>  (cuerpo.docx -> con-portada.docx)
"""
import copy, os, sys
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
W = sys.argv[1]
REPO = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
d = Document(f'{W}/cuerpo.docx')
body = d.element.body
first = body[0]
MARRON = RGBColor(0x4a,0x24,0x16); AMBAR = RGBColor(0xb0,0x63,0x14); GRIS = RGBColor(0x6b,0x55,0x48)
def par(texto='', size=11, color=None, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, after=6):
    p = d.add_paragraph(); p.alignment = align; p.paragraph_format.space_after = Pt(after)
    if texto:
        r = p.add_run(texto); r.font.size = Pt(size); r.bold = bold; r.font.name = 'Liberation Sans'
        if color is not None: r.font.color.rgb = color
    first.addprevious(p._p); return p
par(after=40)
p = par(after=20); p.add_run().add_picture(f'{REPO}/05-erp-odoo/config/logo-almichain.png', width=Cm(6.5))
par('Reto 1 · Almi-Chain', 30, MARRON, True, after=8)
par('Documentación del proyecto', 16, AMBAR, False, after=30)
par('ERP Odoo y TPV · Base de datos con RAID · Web en AWS · Recreativa con MAME · Arcade stick · Red', 11, GRIS, after=50)
par('Alberto · Oier · Julen · Gontzal · Richar', 12, MARRON, True, after=6)
par('1.º ASIR · Almi · Curso 2026-2027', 11, GRIS, after=6)
par('Bilbao, octubre-noviembre de 2026', 11, GRIS, after=6)
# (sin salto de página aquí: el índice ya empieza en página nueva)
# salto de página después del índice
sdt = body.find(qn('w:sdt'))
if sdt is not None:
    p = OxmlElement('w:p'); r = OxmlElement('w:r'); br = OxmlElement('w:br'); br.set(qn('w:type'),'page'); r.append(br); p.append(r)
    sdt.addnext(p)
# el primer capítulo no necesita otro salto (viene del índice): quitar page_break del primer Heading1 se hace por estilo; aceptable
# Pie: número de página; cabecera con el nombre del proyecto
sec = d.sections[0]
sec.different_first_page_header_footer = True
fp = sec.footer.paragraphs[0]; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
def campo(par, instr):
    r = par.add_run(); f1 = OxmlElement('w:fldChar'); f1.set(qn('w:fldCharType'),'begin'); r._r.append(f1)
    r2 = par.add_run(); it = OxmlElement('w:instrText'); it.set(qn('xml:space'),'preserve'); it.text = instr; r2._r.append(it)
    r3 = par.add_run(); f2 = OxmlElement('w:fldChar'); f2.set(qn('w:fldCharType'),'end'); r3._r.append(f2)
campo(fp, 'PAGE')
for r in fp.runs: r.font.size = Pt(9); r.font.color.rgb = GRIS
hp = sec.header.paragraphs[0]; hp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = hp.add_run('Reto 1 · Almi-Chain · Documentación'); r.font.size = Pt(8.5); r.font.color.rgb = GRIS; r.font.name='Liberation Sans'
# actualizar campos al abrir en Word
st = d.settings.element; u = OxmlElement('w:updateFields'); u.set(qn('w:val'),'true'); st.append(u)
d.save(f'{W}/con-portada.docx'); print('portada ok')
